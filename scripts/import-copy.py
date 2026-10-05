#!/usr/bin/env python3
"""Convert Darsh's developer handoff (.docx) + SEO implementation map (.xlsx) into content/pages/*.json.

Usage: python3 scripts/import-copy.py <Developer_Handoff.docx> <SEO_Implementation_Map.xlsx>

Format (handoff 2026-10-04): every page starts with a "Page ID" paragraph ("W02  /  Investment
Management"); the workbook's Final SEO Sitemap gives each Content ID its URL, parent and meta.
Parts I (core pages) and II (Insights) are copy; Parts III/IV, "Implementation" and "Small"
paragraphs are instructions and are never published. Word hyperlinks to p_<ID> bookmarks become
links to that page's URL.

Every word on the site comes from the doc. This script only decides structure (hero, sections,
cards, lists, FAQs, links). Anything it can't place is printed as a WARN line.
"""
import json, re, sys, pathlib
import docx, openpyxl
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "content" / "pages"
SKIP_STYLES = {"Implementation", "Small", "Title"}
STEP_TITLES = re.compile(r"process|how .* works?|what to expect|how we work", re.I)
PILLAR = re.compile(r"^([A-Z][\w &()'-]+?) — (.+)$")
BAND = re.compile(r"^[A-Z0-9 +&]{12,} — .+$")


def warn(page, msg):
    print(f"WARN [{page}] {msg}")


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


# ---------------------------------------------------------------- workbook

def read_sitemap(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Final SEO Sitemap"]
    head = [c.value for c in ws[1]]
    rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(head, r))
        if row.get("Content ID"):
            rows[row["Content ID"]] = row
    by_name = {r["Final Page Name"]: r for r in rows.values()}
    sources = {}
    if "Source Register" in wb.sheetnames:
        sh = wb["Source Register"]
        sh_head = [c.value for c in sh[1]]
        for r in sh.iter_rows(min_row=2, values_only=True):
            s = dict(zip(sh_head, r))
            if s.get("URL"):
                for cid in re.findall(r"[A-Z]{1,2}\d{2}", s.get("Related Content") or ""):
                    sources.setdefault(cid, []).append({"label": s["Authority / Title"], "href": s["URL"]})
    return rows, by_name, sources


# ---------------------------------------------------------------- docx

def para_segments(p, doc, routes):
    """[(text, href|None), ...] with Word hyperlinks resolved to site routes or external URLs."""
    segs = []
    for child in p._p.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "r":
            t = "".join(x.text or "" for x in child.iter(qn("w:t")))
            if t:
                segs.append([t, None])
        elif tag == "hyperlink":
            t = "".join(x.text or "" for x in child.iter(qn("w:t")))
            anchor, rid = child.get(qn("w:anchor")), child.get(qn("r:id"))
            href = None
            if anchor and anchor.startswith("p_"):
                href = routes.get(anchor[2:])
            elif rid:
                href = doc.part.rels[rid].target_ref
            if t:
                segs.append([t, href])
    merged = []
    for t, h in segs:
        if merged and merged[-1][1] == h:
            merged[-1][0] += t
        else:
            merged.append([t, h])
    return merged


def read_doc(path, routes):
    d = docx.Document(path)
    pages, cur = [], None
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(el, d)
            t, s = p.text.strip(), p.style.name
            if not t:
                continue
            if s == "Heading 1" and t.startswith("PART III"):
                break
            if s == "Page ID":
                m = re.match(r"^([A-Z]{1,2}\d{2})\s+/", t)
                cur = {"id": m.group(1), "blocks": []} if m else None
                if cur:
                    pages.append(cur)
                continue
            if cur is None or s in SKIP_STYLES:
                continue
            if s == "Heading 1" and t.startswith("PART "):
                cur = None
                continue
            typ = {"Heading 1": "h1", "Heading 2": "h2", "Heading 3": "h3", "List Bullet": "li",
                   "Subtitle": "h2", "CTA": "cta"}.get(s, "p")
            cur["blocks"].append({"t": typ, "x": t, "segs": para_segments(p, d, routes)})
        elif tag == "tbl" and cur is not None:
            rows = [[c.text.strip() for c in r.cells] for r in Table(el, d).rows]
            cur["blocks"].append({"t": "table", "rows": rows})
    return pages


# ---------------------------------------------------------------- structure

def linked(b):
    """The block is one whole link (e.g. a service name in a list)."""
    segs = [s for s in b.get("segs", []) if s[0].strip()]
    return len(segs) == 1 and bool(segs[0][1]) and segs[0][0].strip() == b["x"].strip()


def para(b):
    out = {"type": "p", "text": b["x"]}
    if any(h for _, h in b.get("segs", [])):
        out["segs"] = b["segs"]
    return out


def button(b, page):
    label = b["x"].strip()
    if re.match(r"schedule", label, re.I):
        return {"type": "button", "label": label, "action": "schedule"}
    href = next((h for _, h in b.get("segs", []) if h), None)
    return {"type": "button", "label": label, "href": href}


def build_blocks(raw, page, article=False, titles=None, hero=False):
    out, i = [], 0
    while i < len(raw):
        b = raw[i]
        t, x = b["t"], b.get("x", "")
        if t == "cta":
            out.append(button(b, page))
            i += 1
            continue
        if t == "table":
            if b["rows"] and b["rows"][0][0] != "Meta Title":
                out.append({"type": "table", "rows": b["rows"]})
            i += 1
            continue
        if t == "li":
            run = []
            while i < len(raw) and raw[i]["t"] == "li":
                run.append(raw[i])
                i += 1
            items = [r["x"] for r in run]
            pills = [PILLAR.match(it) for it in items]
            if all(pills) and len(items) > 1 and not article:
                out.append({"type": "pillars", "items": [
                    {"title": m.group(1).title() if m.group(1).isupper() else m.group(1),
                     "body": [{"type": "p", "text": m.group(2)}]} for m in pills]})
            elif all(linked(r) for r in run):
                out.append({"type": "links", "items": [{"label": r["x"], "href": r["segs"][0][1]} for r in run]})
            else:
                lst = {"type": "list", "items": items}
                if any(any(h for _, h in r.get("segs", [])) for r in run):
                    lst["segs"] = [r["segs"] for r in run]
                out.append(lst)
            continue
        if t == "h3" and article:
            out.append({"type": "h3", "text": x})
            i += 1
            continue
        if t == "h3":
            items = []
            while i < len(raw) and raw[i]["t"] == "h3":
                title = raw[i]["x"]
                i += 1
                body_raw = []
                while i < len(raw) and raw[i]["t"] not in ("h3", "cta", "table"):
                    body_raw.append(raw[i])
                    i += 1
                body = build_blocks(body_raw, page) if body_raw else []
                if title.endswith("?"):
                    if not body:
                        warn(page, f"FAQ with no answer dropped: '{title}'")
                        continue
                    items.append({"q": title, "a": body})
                else:
                    items.append({"title": title, "body": body})
            for it in items:  # keep order; split into faq / card runs
                kind = "faq" if "q" in it else "cards"
                if out and out[-1]["type"] == kind and kind == "faq":
                    out[-1]["items"].append(it)
                elif out and out[-1]["type"] == kind and kind == "cards" and "link" not in out[-1]["items"][-1]:
                    out[-1]["items"].append(it)
                else:
                    out.append({"type": kind, "items": [it]})
            continue
        if t == "p":
            # Featured Insights: a run of lines that are each one link to an article
            if linked(b) and titles is not None:
                items, j = [], i
                while j < len(raw) and raw[j]["t"] == "p" and linked(raw[j]):
                    href = raw[j]["segs"][0][1]
                    items.append({"label": raw[j]["x"], "href": href, "text": titles.get(href, "")})
                    j += 1
                if len(items) >= 2 and all(it["href"].startswith("/insights/") for it in items):
                    out.append({"type": "articles", "items": items})
                    i = j
                    continue
            if BAND.match(x) or (x.upper() == x and (x.count("→") >= 2 or x.count(" + ") >= 2)):
                out.append({"type": "band", "text": x})
                i += 1
                continue
            # a run of short label lines ("Independent Fee-Only Advisor", "450+ data points ...") -> checklist
            short = lambda r: r["t"] == "p" and len(r["x"]) <= 90 and not r["x"].endswith((".", "?", ":")) \
                and " | " not in r["x"] and " • " not in r["x"] and not linked(r) and not BAND.match(r["x"]) \
            and not (r["x"].upper() == r["x"] and " + " in r["x"])
            if not article and short(b):
                j = i
                while j < len(raw) and short(raw[j]):
                    j += 1
                if j - i >= 2 and (hero or j - i == 2):
                    out.append({"type": "chips", "items": [r["x"] for r in raw[i:j]]})
                    i = j
                    continue
                if j - i >= 3:
                    out.append({"type": "list", "items": [r["x"] for r in raw[i:j]]})
                    i = j
                    continue
            if False:
                out.append({"type": "band", "text": x})
                i += 1
                continue
            sep = " | " if " | " in x else " • " if " • " in x else None
            if sep and all(0 < len(s) <= 60 for s in x.split(sep)):
                out.append({"type": "chips", "items": [s.strip() for s in x.split(sep)]})
                i += 1
                continue
            if not article and x.endswith(":") and i + 1 < len(raw) and raw[i + 1]["t"] == "p" and len(raw[i + 1]["x"]) <= 140:
                items, j = [], i + 1
                while j < len(raw) and raw[j]["t"] == "p" and len(raw[j]["x"]) <= 140 \
                        and not raw[j]["x"].endswith(":") and not linked(raw[j]):
                    items.append(raw[j]["x"])
                    j += 1
                if len(items) >= 2:
                    out.append({"type": "list", "lead": x, "items": items})
                    i = j
                    continue
            m = PILLAR.match(x)
            if m and m.group(1).isupper() and not article:
                items, j = [], i
                while j < len(raw) and raw[j]["t"] == "p" and PILLAR.match(raw[j]["x"]) \
                        and PILLAR.match(raw[j]["x"]).group(1).isupper():
                    mm = PILLAR.match(raw[j]["x"])
                    items.append({"title": mm.group(1).title(), "body": [{"type": "p", "text": mm.group(2)}]})
                    j += 1
                out.append({"type": "pillars", "items": items})
                i = j
                continue
            out.append(para(b))
            i += 1
            continue
        warn(page, f"unplaced block {t}: {x[:60]}")
        i += 1
    return merge_card_links(out)


# A few related-service cards (R10-R12) have an unlinked "Learn More" button in the doc; the card
# title names the page, so link by title. Content IDs, resolved to URLs through the workbook.
CARD_TITLE_LINKS = {
    "erisa fiduciary support": "R06", "fiduciary support": "R06",
    "investment analysis": "R15", "401(k) investment analysis": "R15",
    "evaluation & benchmarking": "R11", "401(k) fee benchmarking": "R12",
    "recordkeeper evaluation & rfp consulting": "R13", "retirement plan design": "R21",
    "virtual fiduciary file": "R10",
}
ROUTES = {}


def merge_card_links(blocks):
    """[cards, button] -> the button becomes the last card's link; consecutive linked cards merge."""
    out = []
    for b in blocks:
        prev = out[-1] if out else None
        if b["type"] == "button" and b.get("action") != "schedule" and prev and prev["type"] == "cards" \
                and "link" not in prev["items"][-1]:
            href = b.get("href") or ROUTES.get(CARD_TITLE_LINKS.get(prev["items"][-1]["title"].lower(), ""))
            if href:
                prev["items"][-1]["link"] = {"label": b["label"], "href": href}
                continue
        if b["type"] == "cards" and prev and prev["type"] == "cards" and "link" in prev["items"][-1]:
            prev["items"].extend(b["items"])
            continue
        out.append(b)
    return out


def convert(page, article, titles):
    pid, blocks = page["id"], page["blocks"]
    i = next(k for k, b in enumerate(blocks) if b["t"] == "h1")
    hero = {"title": blocks[i]["x"]}
    i += 1
    if i < len(blocks) and blocks[i]["t"] == "h2":
        hero["subtitle"] = blocks[i]["x"]
        i += 1
    j = i
    while j < len(blocks) and blocks[j]["t"] != "h2":
        j += 1
    hero["blocks"] = build_blocks(blocks[i:j], pid, article, hero=True)
    for b in hero["blocks"]:  # a hero's non-schedule CTA is the secondary action
        if b["type"] == "button" and b.get("action") != "schedule":
            b["secondary"] = True
    sections = []
    while j < len(blocks):
        title = blocks[j]["x"]
        k = j + 1
        while k < len(blocks) and blocks[k]["t"] != "h2":
            k += 1
        sec = {"id": slug(title), "title": title,
               "blocks": build_blocks(blocks[j + 1:k], pid, article, titles)}
        if sec["blocks"] and all(b["type"] in ("faq", "button") for b in sec["blocks"]) \
                and any(b["type"] == "faq" for b in sec["blocks"]):
            sec["layout"] = "faq"
        elif not article and STEP_TITLES.search(title) and any(b["type"] in ("cards", "pillars") for b in sec["blocks"]):
            sec["layout"] = "steps"
        sections.append(sec)
        j = k
    if sections and any(b["type"] == "button" and b.get("action") == "schedule" for b in sections[-1]["blocks"]):
        sections[-1]["layout"] = "cta"
    return hero, sections


def main(doc_path, xlsx_path):
    rows, by_name, sources = read_sitemap(xlsx_path)
    routes = {cid: r["Recommended URL"] for cid, r in rows.items()}
    ROUTES.update(routes)
    pages = read_doc(doc_path, routes)
    missing = set(rows) - {p["id"] for p in pages}
    if missing:
        warn("workbook", f"no copy in the doc for {sorted(missing)}")

    titles = {}  # article subtitles, for the Insights hub cards
    for p in pages:
        if p["id"].startswith("IN"):
            bl = p["blocks"]
            k = next(n for n, b in enumerate(bl) if b["t"] == "h1")
            if k + 1 < len(bl) and bl[k + 1]["t"] == "h2":
                titles[routes[p["id"]]] = bl[k + 1]["x"]

    OUT.mkdir(parents=True, exist_ok=True)
    for f in OUT.glob("*.json"):
        f.unlink()
    for page in pages:
        row = rows.get(page["id"])
        if not row:
            warn(page["id"], "not in the workbook; skipped")
            continue
        article = page["id"].startswith("IN")
        hero, sections = convert(page, article, titles)
        parent = by_name.get(row.get("Parent Page") or "")
        data = {
            "id": page["id"],
            "route": row["Recommended URL"],
            "name": row["Final Page Name"],
            "kind": "article" if article else "page",
            "parentRoute": parent["Recommended URL"] if parent else None,
            "meta": {"title": row.get("Meta Title"), "description": row.get("Meta Description"), "schema": row.get("Schema")},
            "hero": hero,
            "sections": sections,
        }
        for sec in [hero] + sections:
            for b in sec["blocks"]:
                if b["type"] == "button" and b.get("action") != "schedule" and not b.get("href"):
                    warn(page["id"], f"button without a link: '{b['label']}'")
        if article and sources.get(page["id"]):
            data["sources"] = sources[page["id"]]
        fname = slug(row["Recommended URL"]) or "home"
        (OUT / f"{fname}.json").write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n")
        print(f"ok {page['id']} {row['Recommended URL']} ({len(sections)} sections)")

    files = sorted(f.stem for f in OUT.glob("*.json"))
    lines = ["// Generated by scripts/import-copy.py — do not edit by hand.", 'import type { PageContent } from "@/lib/content";']
    for f in files:
        lines.append(f'import p_{f.replace("-", "_")} from "./pages/{f}.json";')
    lines.append("\nexport const pages: PageContent[] = [")
    for f in files:
        lines.append(f"  p_{f.replace('-', '_')} as PageContent,")
    lines.append("];\n")
    (ROOT / "content" / "index.ts").write_text("\n".join(lines))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
